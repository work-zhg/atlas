"""DOCX / PPTX → DocIR（doc/knowledge-enterprise-design.html §4.2）。

DOCX  按文档顺序遍历段落与表格；标题按样式（Heading N / 标题 N）定层级；列表按样式或编号属性；
      段落里内嵌的图片（a:blip）抽出来，紧跟在所在段落之后
PPTX  一页一个小节：标题占位符 → 一级标题；文本框按位置（上 → 下、左 → 右）排；
      段落有缩进级别或项目符号的是列表项；表格、图片、备注（标注「备注」）各成一块；page = 页码
"""

from __future__ import annotations

import io
import re
from typing import Any

from ..docir import Block, DocIR

__all__ = ["parse_docx", "parse_pptx"]

_HEADING_STYLE = re.compile(r"^(?:Heading|标题)\s*(\d)", re.I)
_LIST_STYLE = re.compile(r"List|列表", re.I)
_BLIP = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
_R_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
#: 任一边小于它（像素）的图片多半是图标、分隔线，不抽
_MIN_IMAGE_PX = 48


def _too_small(blob: bytes) -> bool:
    """按像素尺寸判断，不按字节数：纯色或简单的大图压缩后也可能只有几百字节。"""
    try:
        from PIL import Image

        with Image.open(io.BytesIO(blob)) as im:
            w, h = im.size
        return w < _MIN_IMAGE_PX or h < _MIN_IMAGE_PX
    except Exception:  # 不认识的格式（EMF / WMF 等）：太小的不要，其余保留
        return len(blob) < 1_024


def _image_mime(content_type: str | None) -> str:
    return content_type if content_type and content_type.startswith("image/") else "image/png"


# ───────────────────────────────────────────── DOCX


def parse_docx(data: bytes, filename: str = "") -> DocIR:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(io.BytesIO(data))
    ir = DocIR(meta={"parser": "docx"})
    rels = document.part.related_parts
    n_img = n_table = n_head = 0

    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, document)
            text = para.text.strip()
            style = (para.style.name or "") if para.style is not None else ""
            if text:
                if m := _HEADING_STYLE.match(style):
                    ir.blocks.append(Block(type="heading", text=text, level=int(m.group(1))))
                    n_head += 1
                elif style == "Title":
                    ir.blocks.append(Block(type="heading", text=text, level=1))
                    n_head += 1
                elif _LIST_STYLE.search(style) or child.find(".//{*}numPr") is not None:
                    ir.blocks.append(Block(type="list_item", text=text))
                else:
                    ir.blocks.append(Block(type="paragraph", text=text))
            for blip in child.iter(_BLIP):
                rid = blip.get(_R_EMBED)
                part = rels.get(rid) if rid else None
                blob = getattr(part, "blob", None)
                if not blob or _too_small(blob):
                    continue
                n_img += 1
                image_id = f"img{n_img}"
                ir.images[image_id] = (blob, _image_mime(getattr(part, "content_type", None)))
                ir.blocks.append(Block(type="image", attrs={"image_id": image_id}))
        elif tag == "tbl":
            table = Table(child, document)
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            rows = [r for r in rows if any(r)]
            if rows:
                n_table += 1
                ir.blocks.append(Block(type="table", table={"rows": rows, "header_rows": 1}))
    _attach_captions(ir.blocks)
    ir.report = {"headings": n_head, "tables": n_table, "images": n_img}
    return ir


_CAPTION = re.compile(r"^(?:图|表|Figure|Fig\.|Table)\s*[\d一二三四五六七八九十]+[-.\d]*", re.I)


def _attach_captions(blocks: list[Block]) -> None:
    """紧邻图片（前一块或后一块）的「图 3-1 …」段落 → 图注，挂到图片上。"""
    for i, b in enumerate(blocks):
        if b.type != "image" or b.attrs.get("caption"):
            continue
        for j in (i + 1, i - 1):
            if 0 <= j < len(blocks):
                cand = blocks[j]
                if (
                    cand.type == "paragraph"
                    and len(cand.text) <= 80
                    and _CAPTION.match(cand.text.strip())
                ):
                    b.attrs["caption"] = cand.text.strip()
                    cand.type = "caption"
                    cand.attrs["attached"] = True
                    break


# ───────────────────────────────────────────── PPTX


def _shape_order(shape: Any) -> tuple[int, int]:
    return (int(shape.top or 0), int(shape.left or 0))


def parse_pptx(data: bytes, filename: str = "") -> DocIR:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

    prs = Presentation(io.BytesIO(data))
    ir = DocIR(meta={"parser": "pptx", "page_count": len(prs.slides)})
    n_img = n_table = 0
    emu_w = int(prs.slide_width or 1)

    for page, slide in enumerate(prs.slides, start=1):
        title_shape = None
        for sh in slide.placeholders:
            if sh.placeholder_format.type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
                title_shape = sh
                break
        title = title_shape.text_frame.text.strip() if title_shape is not None else ""
        ir.blocks.append(Block(type="heading", text=title or f"第 {page} 页", level=1, page=page))

        def walk(shapes: Any) -> list[Any]:
            out = []
            for sh in shapes:
                if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                    out.extend(walk(sh.shapes))
                else:
                    out.append(sh)
            return out

        for sh in sorted(walk(slide.shapes), key=_shape_order):
            if title_shape is not None and sh.shape_id == title_shape.shape_id:
                continue
            bbox = _emu_bbox(sh, emu_w)
            if sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    text = "".join(r.text for r in para.runs).strip()
                    if not text:
                        continue
                    is_list = para.level > 0 or para._p.find(".//{*}buChar") is not None
                    ir.blocks.append(
                        Block(
                            type="list_item" if is_list else "paragraph",
                            text=text,
                            page=page,
                            bbox=bbox,
                            attrs={"indent": para.level} if para.level else {},
                        )
                    )
            elif getattr(sh, "has_table", False) and sh.has_table:
                rows = [[c.text.strip() for c in r.cells] for r in sh.table.rows]
                rows = [r for r in rows if any(r)]
                if rows:
                    n_table += 1
                    ir.blocks.append(
                        Block(
                            type="table",
                            table={"rows": rows, "header_rows": 1},
                            page=page,
                            bbox=bbox,
                        )
                    )
            elif sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                blob = sh.image.blob
                if _too_small(blob):
                    continue
                n_img += 1
                image_id = f"img{n_img}"
                ir.images[image_id] = (blob, _image_mime(sh.image.content_type))
                ir.blocks.append(
                    Block(type="image", page=page, bbox=bbox, attrs={"image_id": image_id})
                )
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                ir.blocks.append(Block(type="paragraph", text=f"备注：{notes}", page=page))
    _attach_captions(ir.blocks)
    ir.report = {"pages": len(prs.slides), "tables": n_table, "images": n_img}
    return ir


def _emu_bbox(shape: Any, slide_width: int) -> list[float] | None:
    """EMU → 以幻灯片宽度为 1000 的相对坐标（前端按比例换算；不同尺寸的幻灯片也一致）。"""
    try:
        k = 1000 / slide_width
        x0, y0 = int(shape.left) * k, int(shape.top) * k
        return [
            round(x0, 1),
            round(y0, 1),
            round(x0 + int(shape.width) * k, 1),
            round(y0 + int(shape.height) * k, 1),
        ]
    except (TypeError, ValueError):
        return None
