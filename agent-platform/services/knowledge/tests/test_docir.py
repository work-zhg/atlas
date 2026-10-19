"""文档理解（doc/knowledge-enterprise-design.html §4）：DocIR、PDF 版面分析、Office 解析、
结构分段的原子块与引用定位、解析预览与图片接口。"""

from __future__ import annotations

import io
import json
from typing import Any

import pytest
from atlas_knowledge.rag import index_processor as ip
from atlas_knowledge.rag.docir import Block, DocIR
from atlas_knowledge.rag.extractor import extract
from atlas_knowledge.rag.parsers import parse
from conftest import run_indexing
from pdf_samples import employee_handbook, scanned

ALICE = {"X-Uid": "alice"}


@pytest.fixture(scope="module")
def handbook() -> DocIR:
    return parse("员工手册.pdf", employee_handbook())


def _texts(ir: DocIR) -> list[str]:
    return [b.text for b in ir.blocks]


# ───────────────────────────────────────────── PDF 版面分析


def test_pdf_strips_headers_footers_and_page_numbers(handbook: DocIR) -> None:
    joined = "\n".join(_texts(handbook))
    assert "内部资料" not in joined and "第 2 页" not in joined
    assert "第 1 页" in handbook.report["header_footer_removed"]


def test_pdf_reading_order_is_left_column_then_right(handbook: DocIR) -> None:
    texts = _texts(handbook)
    left = next(i for i, t in enumerate(texts) if t.startswith("左栏甲"))
    right = next(i for i, t in enumerate(texts) if t.startswith("右栏乙"))
    assert left < right
    # 左栏的正文都排在右栏之前（不是按行左右交错）
    assert all(handbook.blocks[i].bbox[0] < 300 for i in range(left, right))
    assert handbook.report["two_column_pages"] == 1


def test_pdf_heading_levels_from_font_size_and_numbering(handbook: DocIR) -> None:
    heads = [(b.text, b.level) for b in handbook.blocks if b.type == "heading"]
    # 「第二章 考勤」与正文同字号：按「同一编号模式 = 同一级别」认出来，与「第一章 总则」同级
    assert heads == [("员工手册", 1), ("第一章 总则", 2), ("1.1 适用范围", 3), ("第二章 考勤", 2)]


def test_pdf_cross_page_table_is_merged(handbook: DocIR) -> None:
    (table,) = [b for b in handbook.blocks if b.type == "table"]
    rows = table.table["rows"]
    assert rows[0] == ["假期类型", "天数", "审批人"]
    assert [r[0] for r in rows[1:]] == ["年假", "病假", "婚假", "产假"]  # 续页重复的表头没进来
    assert table.attrs["pages"] == [2, 3]


def test_pdf_image_with_caption(handbook: DocIR) -> None:
    (img,) = [b for b in handbook.blocks if b.type == "image"]
    assert img.attrs["caption"] == "图 1 组织架构" and img.page == 3
    data, mime = handbook.images[img.attrs["image_id"]]
    assert mime == "image/png" and data[:4] == b"\x89PNG"
    caption = next(b for b in handbook.blocks if b.type == "caption")
    assert caption.attrs["attached"] is True


def test_scanned_pdf_is_flagged_not_ocrd() -> None:
    report = parse("扫描件.pdf", scanned()).report
    assert report["scanned_pages"] == [1]
    assert any("OCR" in w for w in report["warnings"])


# ───────────────────────────────────────────── Office / 表格


def _png() -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (200, 120), (200, 30, 30)).save(buf, format="PNG")
    return buf.getvalue()


def test_docx_blocks() -> None:
    import docx

    d = docx.Document()
    d.add_heading("报销制度", level=1)
    d.add_paragraph("差旅报销需要上传发票。")
    d.add_paragraph("机票", style="List Bullet")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "项目", "标准"
    t.cell(1, 0).text, t.cell(1, 1).text = "住宿", "500 元"
    d.add_picture(io.BytesIO(_png()))
    d.add_paragraph("图 2 报销流程")
    buf = io.BytesIO()
    d.save(buf)
    ir = parse("制度.docx", buf.getvalue())
    types = [b.type for b in ir.blocks]
    assert types == ["heading", "paragraph", "list_item", "table", "image", "caption"]
    assert ir.blocks[0].level == 1
    assert ir.blocks[4].attrs["caption"] == "图 2 报销流程"


def test_pptx_blocks() -> None:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "季度目标"
    body = slide.placeholders[1].text_frame
    body.text = "收入增长 20%"
    p = body.add_paragraph()
    p.text = "新签客户 50 家"
    p.level = 1
    rows = slide.shapes.add_table(2, 2, Inches(1), Inches(4), Inches(4), Inches(1)).table
    rows.cell(0, 0).text, rows.cell(0, 1).text = "指标", "目标"
    rows.cell(1, 0).text, rows.cell(1, 1).text = "收入", "1 亿"
    slide.shapes.add_picture(io.BytesIO(_png()), Inches(6), Inches(4))
    slide.notes_slide.notes_text_frame.text = "重点讲客户结构"
    buf = io.BytesIO()
    prs.save(buf)
    ir = parse("规划.pptx", buf.getvalue())
    assert ir.blocks[0].type == "heading" and ir.blocks[0].text == "季度目标"
    kinds = {b.type for b in ir.blocks}
    assert {"list_item", "table", "image"} <= kinds
    assert all(b.page == 1 for b in ir.blocks)
    assert ir.blocks[-1].text == "备注：重点讲客户结构"


def test_tabular_page_documents_match_legacy_format() -> None:
    """XLSX / CSV 走 DocIR 后，普通分段的输入与原解析器一字不差（老知识库行为不变）。"""
    data = "名称,价格\n苹果,3\n香蕉,\n".encode()
    (legacy,) = extract("a.csv", data)
    (doc,) = parse("a.csv", data).page_documents()
    assert doc.page_content == legacy.page_content


def test_plain_text_formats_keep_the_text_path() -> None:
    ir = parse("a.md", "# 标题\n\n正文".encode())
    assert ir.is_plain_text and ir.plain_text() == "# 标题\n\n正文"


def test_docir_json_roundtrip(handbook: DocIR) -> None:
    again = DocIR.from_json(json.loads(json.dumps(handbook.to_json())))
    assert [(b.type, b.text, b.page) for b in again.blocks] == [
        (b.type, b.text, b.page) for b in handbook.blocks
    ]


# ───────────────────────────────────────────── 结构分段：原子块与引用定位


def _rules(**structure: Any) -> dict[str, Any]:
    raw = {"parent_mode": "structure"}
    if structure:
        raw["structure"] = structure
    return ip.normalize_process_rule(
        "hierarchical", raw, doc_form="hierarchical_model", max_length=4000
    )[1]


def test_structure_nodes_carry_locator_and_table_header_context(handbook: DocIR) -> None:
    for b in handbook.blocks:
        if b.type == "image":
            b.asset = "kb/ds/doc/assets/img1.png"
    nodes = ip.transform_ir(handbook, doc_form="hierarchical_model", rules=_rules())
    by_title = {n.metadata["title_path"][-1]: n for n in nodes}
    scope = by_title["1.1 适用范围"]
    loc = scope.metadata["source_locator"]
    assert loc["pages"] == [2, 3] and loc["assets"] == ["kb/ds/doc/assets/img1.png"]
    assert all({"page", "bbox"} <= set(box) for box in loc["boxes"])
    assert "kb-asset:kb/ds/doc/assets/img1.png" in scope.page_content
    table_children = [c for c in scope.children if "婚假" in c.page_content]
    assert (
        table_children and "| 假期类型 | 天数 | 审批人 |" in table_children[0].metadata["context"]
    )


def test_table_is_an_atomic_block() -> None:
    """表格放得进一个父块时不被拆开：放不进当前父块的剩余空间就整体移到下一块。"""
    rows = [["名称", "说明"], *[[f"项目{i}", "说明文字" * 3] for i in range(8)]]
    ir = DocIR(
        blocks=[
            Block(type="heading", text="清单", level=1, page=1),
            Block(type="paragraph", text="前言" * 160, page=1),
            Block(type="table", table={"rows": rows, "header_rows": 1}, page=1),
        ]
    )
    nodes = ip.transform_ir(ir, doc_form="hierarchical_model", rules=_rules(max_length=500))
    with_table = [n for n in nodes if "| 名称 | 说明 |" in n.page_content]
    assert len(with_table) == 1
    assert all(f"项目{i}" in with_table[0].page_content for i in range(8))


def test_non_structure_modes_get_page_locator(handbook: DocIR) -> None:
    _, rules = ip.normalize_process_rule("automatic", None, doc_form="text_model", max_length=4000)
    nodes = ip.transform_ir(handbook, doc_form="text_model", rules=rules)
    assert nodes and all(n.metadata["source_locator"]["pages"] for n in nodes)
    assert not any(n.page_content.startswith("#") for n in nodes)


# ───────────────────────────────────────────── 接口：入库 → 解析预览 → 图片 → 分段定位


async def test_pdf_ingest_preview_asset_and_segment_locator(client: Any, rt: Any) -> None:
    r = await client.post(
        "/v1/datasets",
        json={
            "name": "制度库",
            "embedding_model": "fake-embedding",
            "doc_form": "hierarchical_model",
        },
        headers=ALICE,
    )
    assert r.status_code == 201, r.text
    ds = r.json()
    r = await client.post(
        f"/v1/datasets/{ds['id']}/documents",
        files=[("files", ("员工手册.pdf", employee_handbook(), "application/pdf"))],
        data={
            "process_rule": json.dumps(
                {"mode": "hierarchical", "rules": {"parent_mode": "structure"}}
            )
        },
        headers=ALICE,
    )
    assert r.status_code == 201, r.text
    doc = r.json()["documents"][0]
    await run_indexing(rt)

    base = f"/v1/datasets/{ds['id']}/documents/{doc['id']}"
    info = (await client.get(base, headers=ALICE)).json()
    assert info["indexing_status"] == "completed", info
    assert info["parse_report"]["tables"] == 1 and info["parse_report"]["images"] == 1

    preview = (await client.get(f"{base}/docir", headers=ALICE)).json()
    image = next(b for b in preview["docir"]["blocks"] if b["type"] == "image")
    name = image["asset"].rsplit("/", 1)[-1]
    r = await client.get(f"{base}/assets/{name}", headers=ALICE)
    assert r.status_code == 200 and r.headers["content-type"] == "image/png"
    assert (await client.get(f"{base}/assets/..%2Fx.png", headers=ALICE)).status_code == 404

    segs = (await client.get(f"{base}/segments", headers=ALICE)).json()["data"]
    assert segs and all(s["source_locator"]["pages"] for s in segs)
