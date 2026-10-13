"""纯计算部分：切分、清洗、索引处理器、解析器、重排。"""

from __future__ import annotations

import io

import pytest
from atlas_knowledge.errors import Invalid
from atlas_knowledge.rag.cleaner import clean
from atlas_knowledge.rag.extractor import extract
from atlas_knowledge.rag.index_processor import (
    AUTOMATIC_RULES,
    normalize_process_rule,
    transform,
)
from atlas_knowledge.rag.keyword import bm25_document, bm25_query, tokenize
from atlas_knowledge.rag.models import Document
from atlas_knowledge.rag.rerank import reorder
from atlas_knowledge.rag.splitter import FixedRecursiveCharacterTextSplitter


def test_splitter_respects_size_and_overlap() -> None:
    text = "。".join(f"第{i}句话讲的是知识库的分段规则" for i in range(200))
    sp = FixedRecursiveCharacterTextSplitter(chunk_size=100, chunk_overlap=20, fixed_separator="\n")
    chunks = sp.split_text(text)
    assert len(chunks) > 10
    assert all(len(c) <= 100 for c in chunks)
    # 相邻块有重叠
    assert any(chunks[i][-10:] in chunks[i + 1] for i in range(len(chunks) - 1))


def test_splitter_fixed_separator_first() -> None:
    sp = FixedRecursiveCharacterTextSplitter(chunk_size=500, fixed_separator="\\n\\n")
    assert sp.split_text("甲段\n\n乙段\n\n丙段") == ["甲段", "乙段", "丙段"]


def test_splitter_literal_separator_not_regex() -> None:
    # ". " 不能被当成正则「任意字符 + 空格」
    sp = FixedRecursiveCharacterTextSplitter(
        chunk_size=12, fixed_separator="", separators=[". ", ""]
    )
    out = sp.split_text("ab cd ef gh. ij kl mn op")
    assert out[0].startswith("ab cd ef gh")


def test_cleaner_rules() -> None:
    rules = {
        "pre_processing_rules": [
            {"id": "remove_extra_spaces", "enabled": True},
            {"id": "remove_urls_emails", "enabled": True},
        ]
    }
    raw = "a\n\n\n\nb   c\x07 联系 x@y.com 见 https://e.com/x 与 [文档](https://d.com/a)"
    out = clean(raw, rules)
    assert "\n\n\n" not in out
    assert "b c" in out
    assert "\x07" not in out
    assert "x@y.com" not in out and "https://e.com" not in out
    assert "[文档](https://d.com/a)" in out


def test_normalize_rules() -> None:
    mode, rules = normalize_process_rule("automatic", None, doc_form="text_model", max_length=4000)
    assert mode == "automatic" and rules == AUTOMATIC_RULES
    with pytest.raises(Invalid):
        normalize_process_rule(
            "custom",
            {"segmentation": {"separator": "\n", "max_tokens": 10}},
            doc_form="text_model",
            max_length=4000,
        )
    _, h = normalize_process_rule(
        "hierarchical", {"parent_mode": "full_doc"}, doc_form="hierarchical_model", max_length=4000
    )
    assert h["parent_mode"] == "full_doc" and "segmentation" not in h


def test_parent_child_transform() -> None:
    body = "\n\n".join("\n".join(f"第{p}段第{c}行内容" * 3 for c in range(5)) for p in range(3))
    _, rules = normalize_process_rule(
        "hierarchical",
        {
            "parent_mode": "paragraph",
            "segmentation": {"separator": "\\n\\n", "max_tokens": 500},
            "subchunk_segmentation": {"separator": "\\n", "max_tokens": 60},
        },
        doc_form="hierarchical_model",
        max_length=4000,
    )
    parents = transform([Document(body)], doc_form="hierarchical_model", rules=rules)
    assert len(parents) == 3
    assert all(p.children and len(p.children) == 5 for p in parents)

    _, full = normalize_process_rule(
        "hierarchical", {"parent_mode": "full_doc"}, doc_form="hierarchical_model", max_length=4000
    )
    (only,) = transform([Document(body)], doc_form="hierarchical_model", rules=full)
    assert only.children and len(only.children) >= 3


def test_extractors() -> None:
    import docx
    import openpyxl

    d = docx.Document()
    d.add_heading("退款政策", level=1)
    d.add_paragraph("七天内可以无理由退款。")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "项目", "时限"
    t.cell(1, 0).text, t.cell(1, 1).text = "退款", "7天"
    buf = io.BytesIO()
    d.save(buf)
    (doc,) = extract("policy.docx", buf.getvalue())
    assert "# 退款政策" in doc.page_content and "| 退款 | 7天 |" in doc.page_content

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["问题", "答案"])
    ws.append(["怎么退款", "联系客服"])
    buf = io.BytesIO()
    wb.save(buf)
    (sheet,) = extract("faq.xlsx", buf.getvalue())
    assert '"问题":"怎么退款"' in sheet.page_content

    (csv_doc,) = extract("a.csv", "名称,价格\n苹果,3\n".encode())
    assert '"名称":"苹果";"价格":"3"' in csv_doc.page_content

    (html,) = extract(
        "a.html", b"<html><script>x()</script><body><p>\xe6\xad\xa3\xe6\x96\x87</p></body></html>"
    )
    assert html.page_content == "正文"

    (gbk,) = extract("a.txt", "中文编码测试，这是一段比较长的中文文本用于识别编码。".encode("gbk"))
    assert "中文编码" in gbk.page_content

    with pytest.raises(Invalid):
        extract("a.exe", b"x")


def test_bm25_tokens() -> None:
    assert "退款" in tokenize("如何申请退款？")
    idx, val = bm25_document("退款 退款 发票")
    assert len(idx) == 2 and max(val) > min(val)
    q_idx, q_val = bm25_query("退款的流程")
    assert q_val == [1.0] * len(q_idx)


def test_reorder() -> None:
    assert reorder([1, 2, 3, 4, 5]) == [1, 3, 5, 4, 2]
