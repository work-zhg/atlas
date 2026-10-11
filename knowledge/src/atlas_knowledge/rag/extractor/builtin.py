"""内置解析器的实现。都是同步 CPU 活，由 worker 调用。"""

from __future__ import annotations

import csv
import io

from ..models import Document

__all__ = [
    "decode_text",
    "extract_csv",
    "extract_docx",
    "extract_html",
    "extract_pdf",
    "extract_text",
    "extract_xlsx",
]


def decode_text(data: bytes) -> str:
    if data.startswith(b"\xef\xbb\xbf"):
        return data[3:].decode("utf-8", errors="replace")
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        from charset_normalizer import from_bytes

        best = from_bytes(data).best()
        return str(best) if best is not None else data.decode("utf-8", errors="replace")


def extract_text(data: bytes) -> list[Document]:
    return [Document(page_content=decode_text(data))]


def extract_pdf(data: bytes) -> list[Document]:
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(data)
    docs: list[Document] = []
    try:
        for index in range(len(pdf)):
            page = pdf[index]
            textpage = page.get_textpage()
            try:
                text = textpage.get_text_range()
            finally:
                textpage.close()
                page.close()
            # pdfium 用 \r\n 换行
            text = text.replace("\r\n", "\n").replace("\r", "\n")
            docs.append(Document(page_content=text, metadata={"page": index + 1}))
    finally:
        pdf.close()
    return docs


def _md_table(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [[c.replace("|", "\\|").replace("\n", " ").strip() for c in r] for r in rows]
    rows = [r + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "|" + " --- |" * width]
    lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(lines)


def extract_docx(data: bytes) -> list[Document]:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(io.BytesIO(data))
    parts: list[str] = []
    # 按文档顺序遍历段落与表格（document.paragraphs 会丢表格的位置）
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, document)
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name or "") if para.style is not None else ""
            if style.startswith("Heading"):
                level = style.removeprefix("Heading").strip()
                hashes = "#" * (int(level) if level.isdigit() else 1)
                text = f"{hashes} {text}"
            parts.append(text)
        elif tag == "tbl":
            table = Table(child, document)
            rows = [[cell.text for cell in row.cells] for row in table.rows]
            if rows:
                parts.append(_md_table(rows))
    return [Document(page_content="\n\n".join(parts))]


def _row_text(header: list[str], row: list[object]) -> str:
    pairs = []
    for key, value in zip(header, row, strict=False):
        if value is None or str(value).strip() == "":
            continue
        pairs.append(f'"{key}":"{str(value).strip()}"')
    return ";".join(pairs)


def extract_xlsx(data: bytes) -> list[Document]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    docs: list[Document] = []
    try:
        for ws in wb.worksheets:
            rows = ws.iter_rows(values_only=True)
            header_row = next(rows, None)
            if header_row is None:
                continue
            header = [
                str(h).strip() if h is not None else f"列{i + 1}" for i, h in enumerate(header_row)
            ]
            lines = [t for r in rows if (t := _row_text(header, list(r)))]
            if lines:
                # ★ 每行一段：分隔符用 "\n"，行就是天然的分段单位（同 Dify 的 Excel 处理）
                docs.append(Document(page_content="\n".join(lines), metadata={"sheet": ws.title}))
    finally:
        wb.close()
    return docs


def extract_csv(data: bytes) -> list[Document]:
    reader = csv.reader(io.StringIO(decode_text(data)))
    header = next(reader, None)
    if not header:
        return []
    header = [h.strip() or f"列{i + 1}" for i, h in enumerate(header)]
    lines = [t for r in reader if (t := _row_text(header, list(r)))]
    return [Document(page_content="\n".join(lines))]


def extract_html(data: bytes) -> list[Document]:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(decode_text(data), "html.parser")
    for tag in soup(["script", "style", "noscript", "template"]):
        tag.decompose()
    text = soup.get_text("\n")
    lines = [line.strip() for line in text.splitlines()]
    return [Document(page_content="\n".join(line for line in lines if line))]
